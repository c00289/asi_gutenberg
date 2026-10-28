"""Render a Document to Word (.docx) with python-docx.

Produces a native Word document: real heading styles (so the navigation
pane and Word's own TOC work), a cover page, document-control tables, a TOC
field with a cached entry list, header / footer with PAGE and NUMPAGES
fields, shaded table headers with zebra rows, tinted callouts, requirement
cards, shaded code blocks, task lists, embedded figures and core properties.
Word is asked to refresh fields on open so the TOC gets page numbers.
"""
from __future__ import annotations

import os

from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .inline import parse, plain
from .model import Document
from .render_pdf import wants_cover
from .theme import CALLOUTS, FIGURE_PLACEHOLDER, REQUIREMENT, STATUS_COLOURS, THEME, page_size_pt


def render_docx(doc: Document, path: str, theme: dict | None = None, base_dir: str = ".") -> dict:
    """Write ``doc`` to ``path``.  Returns ``{"paragraphs": n}``."""
    th = dict(THEME)
    if theme:
        th.update(theme)
    r = _DocxRenderer(doc, th, base_dir)
    return r.build(path)


# ------------------------------------------------------------ xml helpers --
def _rgb(hexcol: str) -> RGBColor:
    h = hexcol.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(cell, hexcol: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcol.lstrip("#"))
    tcPr.append(shd)


def _cell_borders(cell, **sides):
    """sides: left=("single", size_eighths_pt, "RRGGBB") ..."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            val, sz, col = sides[side]
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), col.lstrip("#"))
            borders.append(el)
    tcPr.append(borders)


def _table_borders(table, hexcol: str, size: int = 4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), hexcol.lstrip("#"))
        borders.append(el)
    tblPr.append(borders)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _cell_margins(table, left=80, right=80, top=50, bottom=50):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _para_shading(paragraph, hexcol: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcol.lstrip("#"))
    pPr.append(shd)


def _para_border_bottom(paragraph, hexcol: str, size: int = 6):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), hexcol.lstrip("#"))
    pbdr.append(el)
    pPr.append(pbdr)


def _field(paragraph, instr: str, cached: str = ""):
    """Append a simple field (PAGE, NUMPAGES ...) to a paragraph."""
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run._r.append(b)
    run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run._r.append(it)
    run = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run._r.append(s)
    run = paragraph.add_run(cached)
    run = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run._r.append(e)
    return run


def _fld_char(paragraph, kind: str):
    run = paragraph.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._r.append(el)


def _update_fields_on_open(document):
    settings = document.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def _bookmark(paragraph, name: str, bid: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _hyperlink(paragraph, url: str, text: str, colour: str, font: str, size: float):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), colour.lstrip("#"))
    rPr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


# ---------------------------------------------------------------- renderer --
class _DocxRenderer:
    def __init__(self, doc: Document, th: dict, base_dir: str):
        self.doc = doc
        self.th = th
        self.base_dir = base_dir
        self.w = WordDocument()
        self.n_par = 0
        self._bid = 100
        self._toc_seen = False
        W, H = page_size_pt(th)
        self.text_w_in = W / 72.0 - 2 * th["margin_in"]

    # ------------------------------------------------------------ styles --
    def _setup(self):
        th, w = self.th, self.w
        sec = w.sections[0]
        W, H = page_size_pt(th)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width = Pt(W)
        sec.page_height = Pt(H)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(th["margin_in"]))
        st = w.styles
        normal = st["Normal"]
        normal.font.name = th["docx_font"]
        normal.font.size = Pt(th["body_pt"] + 0.5)
        normal.font.color.rgb = _rgb(th["ink"])
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), th["docx_font"])
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for lvl, key in ((1, "h1_pt"), (2, "h2_pt"), (3, "h3_pt"), (4, "h4_pt")):
            h = st[f"Heading {lvl}"]
            h.font.name = th["docx_font"]
            h.element.rPr.rFonts.set(qn("w:eastAsia"), th["docx_font"])
            h.font.size = Pt(th[key])
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = _rgb(th["accent"] if lvl <= 2 else th["ink"])
            h.paragraph_format.space_before = Pt(th[key] * (1.1 if lvl == 1 else 0.8))
            h.paragraph_format.space_after = Pt(th[key] * 0.4)
            h.paragraph_format.keep_with_next = True
        title = st["Title"]
        title.font.name = th["docx_font"]
        title.font.size = Pt(th["title_pt"])
        title.font.bold = True
        title.font.color.rgb = _rgb(th["ink"])
        for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number"):
            try:
                st[name].font.name = th["docx_font"]
            except KeyError:
                pass
        # core properties
        m = self.doc.meta
        cp = w.core_properties
        cp.title = m.title
        cp.subject = m.kind_label
        cp.author = m.author or m.organisation
        cp.keywords = ", ".join(m.keywords)
        cp.comments = " · ".join(x for x in (m.doc_id, f"Rev {m.revision}", m.status) if x)
        cp.category = m.kind_label
        cp.version = m.revision

    # ------------------------------------------------------------- build --
    def build(self, path: str) -> dict:
        self._setup()
        if wants_cover(self.doc):
            self._cover()
        else:
            self._title_only()
        self._header_footer()
        for b in self.doc.front + self.doc.blocks:
            fn = getattr(self, f"_b_{b.kind}", None)
            if fn:
                fn(b)
        if self._toc_seen:
            _update_fields_on_open(self.w)
        self.w.save(path)
        return {"paragraphs": self.n_par}

    # ------------------------------------------------------------- runs ---
    def _runs(self, paragraph, text: str, size: float | None = None, colour: str | None = None,
              bold: bool | None = None, italic: bool | None = None):
        th = self.th
        for s in parse(text):
            if s.br:
                paragraph.add_run().add_break(WD_BREAK.LINE)
                continue
            if s.href:
                _hyperlink(paragraph, s.href, s.text, th["accent"], th["docx_font"], size or th["body_pt"])
                continue
            r = paragraph.add_run(s.text)
            if s.bold or bold:
                r.bold = True
            if s.italic or italic:
                r.italic = True
            if s.strike:
                r.font.strike = True
            if s.code:
                r.font.name = th["docx_mono"]
                r._element.rPr.rFonts.set(qn("w:eastAsia"), th["docx_mono"])
                r.font.size = Pt((size or th["body_pt"]) - 0.5)
                rPr = r._element.get_or_add_rPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), th["code_bg"].lstrip("#"))
                rPr.append(shd)
            elif size:
                r.font.size = Pt(size)
            if colour:
                r.font.color.rgb = _rgb(colour)
        self.n_par += 1
        return paragraph

    def _p(self, text: str = "", style: str | None = None, **kw):
        p = self.w.add_paragraph(style=style)
        if text:
            self._runs(p, text, **kw)
        return p

    # ------------------------------------------------------------- cover --
    def _cover(self):
        th, m, w = self.th, self.doc.meta, self.w
        for _ in range(5):
            w.add_paragraph()
        kind = m.kind_label.upper()
        if m.project:
            kind = f"{m.project.upper()} · {kind}"
        p = w.add_paragraph()
        r = p.add_run(kind)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = _rgb(th["accent"])
        p = w.add_paragraph(style="Title")
        p.add_run(m.title)
        if m.subtitle:
            p = w.add_paragraph()
            r = p.add_run(m.subtitle)
            r.font.size = Pt(14)
            r.font.color.rgb = _rgb(th["muted"])
        p = w.add_paragraph()
        _para_border_bottom(p, th["accent"], 8)
        rows = [(k, v) for k, v in (("Document ID", m.doc_id), ("Revision", m.revision),
                                    ("Status", m.status), ("Date", m.date), ("Author", m.author),
                                    ("Project", m.project), ("Organisation", m.organisation),
                                    ("Classification", m.classification)) if v]
        if rows:
            t = w.add_table(rows=0, cols=2)
            _no_borders(t)
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            for k, v in rows:
                cells = t.add_row().cells
                cells[0].width = Inches(1.5)
                cells[1].width = Inches(self.text_w_in - 1.5)
                pk = cells[0].paragraphs[0]
                rk = pk.add_run(k)
                rk.bold = True
                rk.font.size = Pt(th["small_pt"] + 0.5)
                rk.font.color.rgb = _rgb(th["muted"])
                pv = cells[1].paragraphs[0]
                rv = pv.add_run(v)
                if k == "Status":
                    rv.bold = True
                    rv.font.color.rgb = _rgb(STATUS_COLOURS.get(v, th["ink"]))
        if m.approvers:
            w.add_paragraph()
            p = w.add_paragraph()
            r = p.add_run("APPROVALS")
            r.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = _rgb(th["accent"])
            t = w.add_table(rows=0, cols=2)
            _no_borders(t)
            for role, name in m.approvers:
                cells = t.add_row().cells
                cells[0].width = Inches(1.5)
                cells[1].width = Inches(self.text_w_in - 1.5)
                rk = cells[0].paragraphs[0].add_run(role)
                rk.bold = True
                rk.font.size = Pt(th["small_pt"] + 0.5)
                rk.font.color.rgb = _rgb(th["muted"])
                cells[1].paragraphs[0].add_run(name)
        w.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _title_only(self):
        m = self.doc.meta
        p = self.w.add_paragraph(style="Title")
        p.add_run(m.title)
        if m.subtitle:
            p = self.w.add_paragraph()
            r = p.add_run(m.subtitle)
            r.font.size = Pt(14)
            r.font.color.rgb = _rgb(self.th["muted"])

    def _header_footer(self):
        th, m = self.th, self.doc.meta
        sec = self.w.sections[0]
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.text = ""
        r = hp.add_run(m.title)
        r.font.size = Pt(8)
        r.font.color.rgb = _rgb(th["muted"])
        right = " · ".join(x for x in (m.doc_id, m.kind_label) if x)
        if right:
            hp.add_run("\t\t")
            r = hp.add_run(right)
            r.font.size = Pt(8)
            r.font.color.rgb = _rgb(th["muted"])
        _para_border_bottom(hp, th["rule"], 4)
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        left = " · ".join(x for x in (m.organisation, m.classification, f"Rev {m.revision}" if m.revision else "") if x)
        r = fp.add_run(left)
        r.font.size = Pt(8)
        r.font.color.rgb = _rgb(th["muted"])
        fp.add_run("\t\t")
        r = fp.add_run("Page ")
        r.font.size = Pt(8)
        r.font.color.rgb = _rgb(th["muted"])
        _field(fp, "PAGE", "1")
        r = fp.add_run(" of ")
        r.font.size = Pt(8)
        r.font.color.rgb = _rgb(th["muted"])
        _field(fp, "NUMPAGES", "1")
        for run in fp.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = _rgb(th["muted"])
        if m.status == "Draft":
            # a "DRAFT" tag in the first-page footer keeps the status visible when printed
            f1 = sec.first_page_footer.paragraphs[0]
            f1.text = ""
            r = f1.add_run("DRAFT — not for release")
            r.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = _rgb(STATUS_COLOURS["Draft"])

    # ------------------------------------------------------------ blocks --
    def _b_heading(self, b):
        lvl = min(b.level, 4)
        p = self.w.add_paragraph(style=f"Heading {lvl}")
        self._runs(p, b.display)
        self._bid += 1
        _bookmark(p, f"_h{self._bid}", self._bid)
        if lvl == 1:
            _para_border_bottom(p, self.th["accent"], 6)

    def _b_paragraph(self, b):
        self._p(b.text)

    def _b_raw(self, b):
        self._p(b.text)

    def _b_quote(self, b):
        p = self.w.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        el = OxmlElement("w:left")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "18")
        el.set(qn("w:space"), "8")
        el.set(qn("w:color"), self.th["quote_bar"].lstrip("#"))
        pbdr.append(el)
        pPr.append(pbdr)
        self._runs(p, b.text.replace("\n", "<br>"), italic=True, colour=self.th["muted"])

    def _b_rule(self, b):
        p = self.w.add_paragraph()
        _para_border_bottom(p, self.th["rule"], 6)

    def _b_pagebreak(self, b):
        self.w.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _b_toc(self, b):
        th = self.th
        self._toc_seen = True
        p = self.w.add_paragraph(style="Heading 1")
        p.add_run("Contents")
        _para_border_bottom(p, th["accent"], 6)
        # skip front-matter headings that precede the TOC marker
        seq = self.doc.front + self.doc.blocks
        idx = next((i for i, x in enumerate(seq) if x.kind == "toc"), -1)
        entries = [h for i, h in enumerate(seq) if i > idx and h.kind == "heading" and h.level <= b.depth]
        # complex field spanning several paragraphs: begin / instr / separate ... end
        p = self.w.add_paragraph()
        _fld_char(p, "begin")
        run = p.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f' TOC \\o "1-{b.depth}" \\h \\z \\u '
        run._r.append(it)
        _fld_char(p, "separate")
        first = True
        for h in entries:
            if not first:
                p = self.w.add_paragraph()
            first = False
            p.paragraph_format.left_indent = Inches(0.25 * (h.level - 1))
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(plain(h.display))
            r.font.size = Pt(th["body_pt"] + 0.5)
            if h.level == 1:
                r.bold = True
        if first:
            p = self.w.add_paragraph()
        _fld_char(p, "end")
        note = self.w.add_paragraph()
        r = note.add_run("Right-click the table and choose Update Field to refresh page numbers.")
        r.italic = True
        r.font.size = Pt(th["small_pt"])
        r.font.color.rgb = _rgb(th["muted"])

    def _b_code(self, b):
        th = self.th
        t = self.w.add_table(rows=1, cols=1)
        _table_borders(t, th["rule"], 4)
        _cell_margins(t, 120, 120, 80, 80)
        cell = t.rows[0].cells[0]
        _shade(cell, th["code_bg"])
        lines = b.text.split("\n") or [""]
        p = cell.paragraphs[0]
        for i, ln in enumerate(lines):
            if i:
                p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(ln)
            r.font.name = th["docx_mono"]
            r._element.rPr.rFonts.set(qn("w:eastAsia"), th["docx_mono"])
            r.font.size = Pt(th["code_pt"] + 0.5)
            r.font.color.rgb = _rgb(th["code_ink"])
        self.w.add_paragraph().paragraph_format.space_after = Pt(2)

    def _b_list(self, b):
        self._list_items(b.items, b.ordered, 0)

    def _list_items(self, items, ordered, depth):
        for i, it in enumerate(items, 1):
            if it.checked is not None:
                p = self.w.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
                p.paragraph_format.first_line_indent = Inches(-0.2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(("☑" if it.checked else "☐") + " ")
                r.font.name = "Segoe UI Symbol"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI Symbol")
                self._runs(p, it.text)
            elif ordered:
                p = self.w.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3 * (depth + 1))
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.paragraph_format.space_after = Pt(2)
                p.add_run(f"{i}.\t")
                self._runs(p, it.text)
            else:
                style = "List Bullet" if depth == 0 else f"List Bullet {min(depth + 1, 3)}"
                p = self.w.add_paragraph(style=style)
                p.paragraph_format.space_after = Pt(2)
                self._runs(p, it.text)
            if it.children:
                self._list_items(it.children, ordered, depth + 1)

    def _col_widths_in(self, b) -> list[float]:
        ncol = len(b.columns)
        weights = []
        for ci in range(ncol):
            lens = [len(plain(b.columns[ci]))] + [len(plain(r[ci])) for r in b.rows]
            longest = max(lens)
            avg = sum(lens) / len(lens)
            weights.append(max(min(longest, 48) * 0.6 + avg * 0.4, 5.0))
        total = sum(weights)
        return [self.text_w_in * wgt / total for wgt in weights]

    def _b_table(self, b):
        th = self.th
        ncol = len(b.columns)
        widths = self._col_widths_in(b)
        t = self.w.add_table(rows=1, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _table_borders(t, th["rule"], 4)
        _cell_margins(t)
        hdr = t.rows[0]
        _repeat_header(hdr)
        align = list(b.align) + ["left"] * (ncol - len(b.align))
        for ci, c in enumerate(b.columns):
            cell = hdr.cells[ci]
            cell.width = Inches(widths[ci])
            _shade(cell, th["accent"])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            self._runs(p, c, size=th["body_pt"], colour="#FFFFFF", bold=True)
            p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                           "right": WD_ALIGN_PARAGRAPH.RIGHT}[align[ci]]
        for ri, r in enumerate(b.rows):
            row = t.add_row()
            for ci, c in enumerate(r):
                cell = row.cells[ci]
                cell.width = Inches(widths[ci])
                if ri % 2 == 1:
                    _shade(cell, th["zebra_b"])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                self._runs(p, c.replace("\n", "<br>"), size=th["body_pt"])
                p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                               "right": WD_ALIGN_PARAGRAPH.RIGHT}[align[ci]]
        self._caption(b.display_caption)

    def _caption(self, text: str):
        p = self.w.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(10)
        if text:
            self._runs(p, text, size=self.th["small_pt"] + 0.5, colour=self.th["muted"], italic=True)

    def _b_callout(self, b):
        label, fill, stroke = CALLOUTS[b.style]
        t = self.w.add_table(rows=1, cols=1)
        _no_borders(t)
        _cell_margins(t, 160, 160, 100, 100)
        cell = t.rows[0].cells[0]
        cell.width = Inches(self.text_w_in)
        _shade(cell, fill)
        _cell_borders(cell, left=("single", 24, stroke))
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(b.title or label)
        r.bold = True
        r.font.color.rgb = _rgb(stroke)
        for para in b.text.split("\n\n"):
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            self._runs(p, para.replace("\n", " "))
        self.w.add_paragraph().paragraph_format.space_after = Pt(2)

    def _b_requirement(self, b):
        th = self.th
        fill, stroke = REQUIREMENT
        t = self.w.add_table(rows=2 if not b.attrs else 3, cols=1)
        _table_borders(t, stroke, 4)
        _cell_margins(t, 140, 140, 70, 70)
        head = t.rows[0].cells[0]
        head.width = Inches(self.text_w_in)
        _shade(head, fill)
        _cell_borders(head, left=("single", 24, stroke))
        p = head.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("«requirement»")
        r.font.size = Pt(th["small_pt"])
        r.font.color.rgb = _rgb(stroke)
        p = head.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{b.req_id} — ")
        r.bold = True
        r.font.color.rgb = _rgb(stroke)
        self._runs(p, b.title, bold=True, colour=stroke)
        body = t.rows[1].cells[0]
        _cell_borders(body, left=("single", 24, stroke))
        p = body.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        self._runs(p, b.text.replace("\n", "<br>"))
        if b.attrs:
            ac = t.rows[2].cells[0]
            _cell_borders(ac, left=("single", 24, stroke))
            p = ac.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            for i, (k, v) in enumerate(b.attrs.items()):
                if i:
                    rr = p.add_run("  ·  ")
                    rr.font.size = Pt(th["small_pt"])
                    rr.font.color.rgb = _rgb(th["muted"])
                rk = p.add_run(f"{k} ")
                rk.bold = True
                rk.font.size = Pt(th["small_pt"])
                rk.font.color.rgb = _rgb(th["muted"])
                rv = p.add_run(v)
                rv.font.size = Pt(th["small_pt"])
                rv.font.color.rgb = _rgb(th["muted"])
        self.w.add_paragraph().paragraph_format.space_after = Pt(2)

    def _b_figure(self, b):
        th = self.th
        path = b.path if os.path.isabs(b.path) else os.path.join(self.base_dir, b.path)
        frac = b.width or 0.9
        placed = False
        if os.path.isfile(path):
            try:
                p = self.w.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Inches(self.text_w_in * frac))
                placed = True
            except Exception:
                placed = False
        if not placed:
            fill, stroke = FIGURE_PLACEHOLDER
            t = self.w.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            _table_borders(t, stroke, 6)
            cell = t.rows[0].cells[0]
            cell.width = Inches(self.text_w_in * 0.7)
            _shade(cell, fill)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(24)
            r = p.add_run("figure placeholder\n")
            r.font.size = Pt(th["small_pt"])
            r.font.color.rgb = _rgb(th["muted"])
            r = p.add_run(b.path)
            r.font.name = th["docx_mono"]
            r.font.size = Pt(th["small_pt"])
            r.font.color.rgb = _rgb(th["muted"])
        self._caption(b.display_caption)
