"""The in-place fill (0.7.0): the Word original with the source's text written into its own runs; the
abridged headers of 0.8.0 (verbatim on a table's first page, short on every later one, a Word pass)."""

from __future__ import annotations

import json
import os
import pathlib
import subprocess
import sys

import pytest

from asi_gutenberg import fill_docx, from_docx, from_markdown, to_markdown
from asi_gutenberg.theme import THEME

REPO = pathlib.Path(__file__).resolve().parent.parent
QF17 = REPO.parent / "asi_standards" / "qf-17-safety-plan"
PROGRAM_FILES = pathlib.Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
WINWORD = PROGRAM_FILES / "Microsoft Office" / "root" / "Office16" / "WINWORD.EXE"
GREEN = THEME["mark_bg"].lstrip("#")  # the mark colour, whatever the theme says it is


def _template(path: pathlib.Path) -> pathlib.Path:
    """A small Word original shaped like a controlled form: a title, control table, a bold prefix before a
    blank, bullets, a landscape section holding a table whose cells have lines and blanks, a footer."""
    import docx
    from docx.enum.section import WD_ORIENT, WD_SECTION

    d = docx.Document()
    d.add_paragraph("Widget Safety Plan Template", style="Title")
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text, t.cell(0, 1).text = "Document ID", "QF-99"
    t.cell(1, 0).text, t.cell(1, 1).text = "Document owner", "__________"
    d.add_paragraph("Purpose", style="Heading 1")
    p = d.add_paragraph()
    p.add_run("Note: ").bold = True
    p.add_run("this plan applies to ___________ program(s), ")
    p.add_run("always.").italic = True
    d.add_paragraph("Perform systematic planning.", style="List Bullet")
    d.add_paragraph("Assign the activities.", style="List Bullet")
    d.add_paragraph("Plan", style="Heading 1")
    wide = d.add_section(WD_SECTION.NEW_PAGE)
    wide.orientation = WD_ORIENT.LANDSCAPE
    wide.page_width, wide.page_height = wide.page_height, wide.page_width
    t2 = d.add_table(rows=3, cols=3)
    t2.style = "Table Grid"
    for c, h in enumerate(("Phase", "Deliverable", "Deadline")):
        t2.cell(0, c).text = h
    t2.cell(1, 0).text, t2.cell(1, 1).text, t2.cell(1, 2).text = "1", "Safety Plan", "__________"
    t2.cell(2, 0).text = "2"
    t2.cell(2, 1).paragraphs[0].text = "Role A"
    t2.cell(2, 1).add_paragraph("Role B")
    back = d.add_section(WD_SECTION.NEW_PAGE)
    back.orientation = WD_ORIENT.PORTRAIT
    back.page_width, back.page_height = back.page_height, back.page_width
    d.add_paragraph("Resolution", style="Heading 1")
    d.add_paragraph("Closing text.")
    d.sections[0].footer.paragraphs[0].text = "Proprietary"
    d.save(str(path))
    return path


def _source(template: pathlib.Path, tmp_path: pathlib.Path) -> str:
    """The template's Markdown source with its blanks named, as the owner names them."""
    src = from_docx(str(template), str(tmp_path / "figures")).document
    src.meta.title = "{{ vehicle }} Safety Plan"
    tables = [b for b in src.blocks if b.kind == "table"]
    tables[0].rows[1][1] = "{{ owner }}"
    para = next(b for b in src.blocks if b.kind == "paragraph")
    para.text = para.text.replace("___________", "the {{ program }}")
    tables[1].rows[0][2] = "{{ deadline }}"
    tables[1].rows[1][1] = "Role A: {{ a }}\nRole B: {{ b }}\nRole C: {{ c }}"
    tables[1].rows[1][2] = "{{ status }}"
    return to_markdown(src)


VALUES = {"vehicle": "Wheel Loader", "owner": "Owner One", "program": "Alpha", "deadline": "2027-01-31",
          "a": "Ann", "b": "Ben", "c": "Cy", "status": "Pending"}


def _texts(cell) -> list[str]:
    return [p.text for p in cell.paragraphs]


SHORT = ["PDP", "Lifecycle", "Deliverable", "Release", "Reference", "Author", "Verification", "Gate reviewer",
         "Duration", "Status"]


def _wide_template(path: pathlib.Path) -> pathlib.Path:
    """A landscape original with one ten-column, 80-row plan table under a wordy, formatted header (bold Arial
    9pt, a second paragraph in one cell) that runs over several pages."""
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    d = docx.Document()
    d.add_paragraph("Wide Plan", style="Title")
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    heads = ["Phase of the Product Development Process", "Safety Lifecycle Phase", "Safety Deliverables (every one)",
             "Work Product Release (location, document and version control)", "QMS/PDP/Platform reference",
             "Author (Designation: Name)", "Verification Independence Reviewer: Name",
             "Gate Reviewer (Designation: Name)", "Expected Duration/Deadline (to be planned)",
             "Status (Pending/Ongoing/Blocked/Done)"]
    t = d.add_table(rows=81, cols=10)
    t.style = "Table Grid"
    for c, h in enumerate(heads):
        run = t.cell(0, c).paragraphs[0].add_run(h)
        run.bold, run.font.size, run.font.name = True, Pt(9), "Arial"
    t.cell(0, 3).add_paragraph("Follow the plan.").runs[0].font.size = Pt(8)
    for r in range(1, 81):
        for c in range(10):
            t.cell(r, c).paragraphs[0].add_run("__________" if (r, c) == (1, 8) else f"r{r}c{c}").font.size = Pt(8)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    d.add_paragraph("After the table.")
    d.save(str(path))
    return path


def _grid(table) -> list[str]:
    from docx.oxml.ns import qn

    return [g.get(qn("w:w")) for g in table._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]


def _is_header(table) -> bool:
    from docx.oxml.ns import qn

    return table.rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None


def test_fill_writes_the_values_into_the_original_and_nothing_else(tmp_path):
    import docx
    from docx.oxml.ns import qn

    template = _template(tmp_path / "widget.docx")
    result = fill_docx(str(template), _source(template, tmp_path), VALUES, str(tmp_path / "out" / "widget.docx"))
    assert result.replaced == 4 and result.unfilled == [] and result.issues == [] and result.pdf is None
    out = docx.Document(result.docx)
    # the title, in its paragraph and in the file's properties
    assert out.paragraphs[0].text == "Wheel Loader Safety Plan" and out.paragraphs[0].style.name == "Title"
    assert out.core_properties.title == "Wheel Loader Safety Plan"
    # the bold prefix and the italic tail keep their runs; only the blank between them changed
    note = next(p for p in out.paragraphs if p.text.startswith("Note:"))
    assert note.text == "Note: this plan applies to the Alpha program(s), always."
    assert note.runs[0].text == "Note: " and note.runs[0].bold
    assert note.runs[-1].text == "always." and note.runs[-1].italic
    assert not any(r.bold or r.italic for r in note.runs[1:-1])
    # the control table, the plan table: text, lines as paragraphs, an empty cell filled, styles kept
    control, plan = out.tables
    assert control.cell(1, 1).text == "Owner One" and control.style.name == "Table Grid"
    assert plan.cell(1, 2).text == "2027-01-31" and plan.cell(2, 2).text == "Pending"
    assert _texts(plan.cell(2, 1)) == ["Role A: Ann", "Role B: Ben", "Role C: Cy"]
    assert plan.cell(2, 1).paragraphs[2].style.name == plan.cell(2, 1).paragraphs[0].style.name
    assert plan.style.name == "Table Grid"
    # every table's first row repeats on each page; the sections and the footer are the original's
    for table in out.tables:
        assert table.rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
    assert [s.orientation.name for s in out.sections] == ["PORTRAIT", "LANDSCAPE", "PORTRAIT"]
    assert out.sections[0].footer.paragraphs[0].text == "Proprietary"
    assert [p.style.name for p in out.paragraphs if p.style.name.startswith("List")] == ["List Bullet", "List Bullet"]
    # read back through the importer, the file says exactly what the filled source says
    again = from_docx(result.docx, str(tmp_path / "figures2")).document
    assert again.title == "Wheel Loader Safety Plan"
    assert [b for b in again.blocks if b.kind == "table"][1].rows[1][1] == "Role A: Ann\nRole B: Ben\nRole C: Cy"


def test_marks_shade_the_values_only(tmp_path):
    import docx
    from docx.oxml.ns import qn

    template = _template(tmp_path / "widget.docx")
    result = fill_docx(str(template), _source(template, tmp_path), VALUES, str(tmp_path / "marked.docx"), mark=True)
    out = docx.Document(result.docx)
    note = next(p for p in out.paragraphs if p.text.startswith("Note:"))

    def fill_of(run) -> str | None:
        shd = run._r.find(qn("w:rPr") + "/" + qn("w:shd")) if run._r.find(qn("w:rPr")) is not None else None
        return shd.get(qn("w:fill")) if shd is not None else None

    shaded = [r.text for r in note.runs if fill_of(r) == GREEN]
    assert shaded == ["Alpha"] and fill_of(note.runs[0]) is None
    plan = out.tables[1]
    assert [r.text for p in plan.cell(2, 1).paragraphs for r in p.runs if fill_of(r) == GREEN] == ["Ann", "Ben", "Cy"]
    plain = fill_docx(str(template), _source(template, tmp_path), VALUES, str(tmp_path / "plain.docx"))
    assert not any(fill_of(r) == GREEN for p in docx.Document(plain.docx).paragraphs for r in p.runs)


def test_unfilled_placeholders_stay_or_refuse(tmp_path):
    import docx

    template = _template(tmp_path / "widget.docx")
    source = _source(template, tmp_path)
    with pytest.raises(ValueError, match="no value for parameter"):
        fill_docx(str(template), source, {"vehicle": "X"}, str(tmp_path / "strict.docx"), strict=True)
    result = fill_docx(str(template), source, {"vehicle": "Wheel Loader"}, str(tmp_path / "loose.docx"))
    assert result.unfilled == ["a", "b", "c", "deadline", "owner", "program", "status"]
    out = docx.Document(result.docx)
    assert out.tables[0].cell(1, 1).text == "{{ owner }}" and out.paragraphs[0].text == "Wheel Loader Safety Plan"


def test_nothing_to_change_changes_nothing(tmp_path):
    template = _template(tmp_path / "widget.docx")
    untouched = to_markdown(from_docx(str(template), str(tmp_path / "figures")).document)
    result = fill_docx(str(template), untouched, {}, str(tmp_path / "same.docx"), repeat_header=False)
    assert result.replaced == 0
    before = from_docx(str(template), str(tmp_path / "f1")).document.to_dict()
    after = from_docx(result.docx, str(tmp_path / "f2")).document.to_dict()
    assert before == after


def test_a_cloned_line_keeps_the_runs_properties(tmp_path):
    """0.7.1: a cell written line by line clones its last paragraph for the extra lines; the clone's runs are
    emptied, not dropped, so the text written there carries their properties (bold, size, font) as the first
    line's does, and an empty paragraph with no run takes its paragraph mark's. Before, both went into a bare
    run and took the paragraph style's look: the extra role lines of the QF-17 came out in the body font."""
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = docx.Document()
    d.add_paragraph("Roles", style="Title")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Role", "Who"
    run = t.cell(1, 0).paragraphs[0].add_run("Role A: ____")
    run.bold, run.font.size, run.font.name = True, Pt(7), "Arial"
    blank = t.cell(1, 1).paragraphs[0]  # no run at all; the paragraph mark says bold
    mark = blank._p.get_or_add_pPr()
    rpr = mark.makeelement(qn("w:rPr"), {})
    rpr.append(rpr.makeelement(qn("w:b"), {}))
    mark.append(rpr)
    template = tmp_path / "roles.docx"
    d.save(str(template))
    src = from_docx(str(template), str(tmp_path / "figures")).document
    [table] = [b for b in src.blocks if b.kind == "table"]
    table.rows[0][0] = "Role A: {{ a }}" + chr(10) + "Role B: {{ b }}" + chr(10) + "Role C: {{ c }}"
    table.rows[0][1] = "{{ who }}"
    result = fill_docx(str(template), to_markdown(src), {"a": "Ann", "b": "Ben", "c": "Cy", "who": "Ann"},
                       str(tmp_path / "filled.docx"))
    out = docx.Document(result.docx).tables[0]
    assert _texts(out.cell(1, 0)) == ["Role A: Ann", "Role B: Ben", "Role C: Cy"]
    for paragraph in out.cell(1, 0).paragraphs:
        for r in paragraph.runs:
            if r.text:
                assert (r.bold, r.font.size, r.font.name) == (True, Pt(7), "Arial"), paragraph.text
    [who] = [r for r in out.cell(1, 1).paragraphs[0].runs if r.text]
    assert who.text == "Ann" and who.bold


def test_a_failed_print_leaves_the_pdf_that_was_there(tmp_path, monkeypatch):
    """0.7.1: Word prints to a working name and the target changes hands only once the print is there, so a
    print that fails (no Word, a COM error) leaves the PDF that was at the target, which asi_kletz counts on:
    its export writes Gutenberg's rendering first and keeps it when Word cannot print. Before, the target was
    removed before Word was asked, and a failed print left no PDF at all."""
    import subprocess

    template = _template(tmp_path / "widget.docx")
    rendered = tmp_path / "printed.pdf"
    rendered.write_bytes(b"%PDF-1.4 the rendering of the source")

    def no_word(*args, **kwargs):
        raise OSError("powershell is not here")

    monkeypatch.setattr(subprocess, "run", no_word)  # what word.py calls, by module attribute
    result = fill_docx(str(template), _source(template, tmp_path), VALUES, str(tmp_path / "printed.docx"), pdf=True)
    assert result.pdf is None and result.issues == ["pdf through Word failed: powershell is not here"]
    assert rendered.read_bytes() == b"%PDF-1.4 the rendering of the source"
    assert not (tmp_path / "printed.printing.pdf").exists()


def test_a_source_out_of_line_with_the_original_is_refused(tmp_path):
    template = _template(tmp_path / "widget.docx")
    src = from_markdown(_source(template, tmp_path))
    src.blocks = [b for b in src.blocks if b.kind != "list"]  # two bullets gone: the units no longer line up
    with pytest.raises(ValueError, match="unit 5: the Word file has a text, the source a table"):
        fill_docx(str(template), src, VALUES, str(tmp_path / "bad.docx"))
    src = from_markdown(_source(template, tmp_path))
    table = [b for b in src.blocks if b.kind == "table"][1]
    table.rows.append(["3", "Extra", "x"])
    with pytest.raises(ValueError, match="rows"):
        fill_docx(str(template), src, VALUES, str(tmp_path / "bad2.docx"))


def test_the_command_line_fills(tmp_path):
    template = _template(tmp_path / "source" / "widget.docx") if (tmp_path / "source").mkdir() is None else None
    md = tmp_path / "widget.md"
    md.write_text(_source(template, tmp_path), encoding="utf-8")
    (tmp_path / "values.json").write_text(json.dumps(VALUES), encoding="utf-8")
    run = subprocess.run([sys.executable, "-m", "asi_gutenberg", "fill", str(md), str(tmp_path / "out" / "w.docx"),
                          "--params-file", str(tmp_path / "values.json"), "--param", "program=Beta"],
                         capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 0, run.stderr
    payload = json.loads(run.stdout)
    assert payload["ok"] and payload["verb"] == "fill" and payload["replaced"] == 4 and payload["pdf"] is None
    assert payload["template"] == str(template.resolve()) and os.path.isfile(payload["docx"])
    import docx
    note = next(p.text for p in docx.Document(payload["docx"]).paragraphs if p.text.startswith("Note"))
    assert "the Beta program" in note
    run = subprocess.run([sys.executable, "-m", "asi_gutenberg", "fill", str(md), str(tmp_path / "x.docx")],
                         capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 1 and "no value for parameter" in json.loads(run.stdout)["error"]
    # 0.8.0: --abridge-headers takes a JSON object; without --pdf the pass cannot run and the payload says so
    (tmp_path / "headings.json").write_text(json.dumps({"Phase": "Ph.", "Deliverable": "Deliv.", "Deadline": "Due"}),
                                            encoding="utf-8")
    run = subprocess.run([sys.executable, "-m", "asi_gutenberg", "fill", str(md), str(tmp_path / "out" / "a.docx"),
                          "--params-file", str(tmp_path / "values.json"), "--abridge-headers",
                          str(tmp_path / "headings.json")], capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 0, run.stderr
    payload = json.loads(run.stdout)
    assert payload["abridged"] == [] and payload["pages"] is None
    assert payload["issues"] == ["abridged headers need Word (pdf=True on Windows with Word installed)"]
    (tmp_path / "headings.json").write_text(json.dumps({"Phase": "Ph."}), encoding="utf-8")
    run = subprocess.run([sys.executable, "-m", "asi_gutenberg", "fill", str(md), str(tmp_path / "out" / "b.docx"),
                          "--params-file", str(tmp_path / "values.json"), "--abridge-headers",
                          str(tmp_path / "headings.json")], capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 1 and "headings without a short form" in json.loads(run.stdout)["error"]


@pytest.mark.skipif(not WINWORD.is_file(), reason="Word is not installed here")
def test_word_prints_the_pdf(tmp_path):
    template = _template(tmp_path / "widget.docx")
    result = fill_docx(str(template), _source(template, tmp_path), VALUES, str(tmp_path / "printed.docx"), pdf=True)
    assert result.issues == [] and result.pdf == str(tmp_path / "printed.pdf")
    assert pathlib.Path(result.pdf).stat().st_size > 1000


@pytest.mark.skipif(not (QF17 / "qf-17-safety-plan.md").is_file(),
                    reason="asi_standards (private) is not checked out beside this repo")
def test_qf_17_fills_whole(tmp_path):
    """The owner's controlled document: every unit lines up, every value lands, nothing else moves."""
    import docx
    from docx.oxml.ns import qn

    from asi_gutenberg import fill, parameters

    source = (QF17 / "qf-17-safety-plan.md").read_text(encoding="utf-8")
    template = next((QF17 / "source").glob("*.docx"))
    doc = from_markdown(source)
    values = {n: f"<{n}>" for n in parameters(doc)}
    values["vehicle_type"] = "Wheel Loader"
    result = fill_docx(str(template), source, values, str(tmp_path / "qf-17.docx"))
    assert result.unfilled == [] and result.replaced >= 3
    out = docx.Document(result.docx)
    assert out.core_properties.title == "Wheel Loader Safety Plan (SP)"
    assert [s.orientation.name for s in out.sections] == ["PORTRAIT", "LANDSCAPE", "PORTRAIT"]
    plan = out.tables[2]
    assert len(plan.rows) == 50 and len(plan.columns) == 10 and plan.style.name == "Grid Table 6 Colorful"
    assert plan.rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
    durations = [plan.cell(r, 8).text for r in range(1, len(plan.rows))]
    statuses = [plan.cell(r, 9).text for r in range(1, len(plan.rows))]
    assert durations[0] == "<row_1_points> points\n<row_1_sprint_end>"
    assert any("<row_48_status>" in s for s in statuses)  # the Status column stacks the Jira key over the status
    # 0.9.0: the deliverable's name is a link to the row's first work product URL, blue, not green
    assert _links_of(plan.cell(1, 2)) == [("<row_1_url>", "Safety Plan", "Hyperlink", "0563C1", BLUE)]
    assert not any(s.get(qn("w:fill")) == GREEN for s in plan.cell(1, 2)._tc.iter(qn("w:shd")))
    assert len([h for r in range(1, len(plan.rows)) for h in _links_of(plan.cell(r, 2))]) >= 49
    # 0.10.0: row 1's url a placeholder and its key a gap: pinkish red, the link kept; row 2 untouched
    gaps = fill_docx(str(template), source, values, str(tmp_path / "qf-17-gaps.docx"), mark=True,
                     warn=["row_1_jira_key"], warn_urls=["<row_1_url>"])
    assert gaps.unfilled == []
    gap_plan = docx.Document(gaps.docx).tables[2]
    assert _links_of(gap_plan.cell(1, 2)) == [("<row_1_url>", "Safety Plan", "Hyperlink", DARK_RED, PINK)]
    assert _links_of(gap_plan.cell(2, 2))[0][3:] == ("0563C1", BLUE)
    status_fills = _fills_of(gap_plan.cell(1, 9))
    assert ("<row_1_jira_key>", PINK) in status_fills and ("<row_1_status>", GREEN) in status_fills
    # read back, the file says exactly what the filled source says (less its marks and links), cell for cell
    from asi_gutenberg.fill_docx import _plain, _source_units

    want = [(kind, _plain(unit) if kind == "text" else [[_plain(c) for c in row] for row in unit])
            for kind, unit in _source_units(fill(doc, values, mark=False))]
    got = _source_units(from_docx(result.docx, str(tmp_path / "figures")).document)
    assert len(want) == len(got)
    for w, g in zip(want, got, strict=True):
        assert w == g


def test_abridged_headers_are_checked_and_need_word(tmp_path, monkeypatch):
    """0.8.0: the mapping is checked before anything is written; a table only partly in it is a typo, not a
    table to skip; without Word (no ``pdf=True``, or no PowerShell at all) the pass cannot run, the result
    says so and the file keeps the verbatim header repeating on every page."""
    import docx

    from asi_gutenberg.fill_docx import ABRIDGE_NEEDS_WORD

    template = _template(tmp_path / "widget.docx")
    source = _source(template, tmp_path)
    short = {"Phase": "Ph.", "Deliverable": "Deliv.", "Deadline": "Due"}
    for bad in ({}, {"Phase": ""}, {"Phase": 1}, {3: "x"}, "Phase=Ph."):
        with pytest.raises(ValueError, match="abridged_headers"):
            fill_docx(str(template), source, VALUES, str(tmp_path / "bad.docx"), abridged_headers=bad)
    with pytest.raises(ValueError, match=r"headings without a short form: table 2 \(Phase\): 'Deadline'"):
        fill_docx(str(template), source, VALUES, str(tmp_path / "bad.docx"),
                  abridged_headers={"Phase": "Ph.", "Deliverable": "Deliv."})
    with pytest.raises(ValueError, match="no table's first row is in the mapping"):
        fill_docx(str(template), source, VALUES, str(tmp_path / "bad.docx"), abridged_headers={"Nowhere": "x"})
    assert not (tmp_path / "bad.docx").exists()

    def verbatim_and_repeating(path: str) -> None:
        out = docx.Document(path)
        assert len(out.tables) == 2
        plan = out.tables[1]
        assert [c.text for c in plan.rows[0].cells] == ["Phase", "Deliverable", "Deadline"] and _is_header(plan)

    result = fill_docx(str(template), source, VALUES, str(tmp_path / "plain.docx"), abridged_headers=short)
    assert result.issues == [ABRIDGE_NEEDS_WORD] and result.abridged == [] and result.pages is None
    verbatim_and_repeating(result.docx)

    def no_word(*args, **kwargs):
        raise OSError("powershell is not here")

    monkeypatch.setattr(subprocess, "run", no_word)
    result = fill_docx(str(template), source, VALUES, str(tmp_path / "printed.docx"), pdf=True, abridged_headers=short)
    assert result.pdf is None and result.abridged == [] and result.pages is None
    assert result.issues == ["pdf through Word failed: powershell is not here", ABRIDGE_NEEDS_WORD]
    verbatim_and_repeating(result.docx)


@pytest.mark.skipif(not WINWORD.is_file(), reason="Word is not installed here")
def test_word_abridges_the_headers_after_page_one(tmp_path):
    """0.8.0, with Word: the table is split where its first page ends; the first part keeps the verbatim
    header, shown once; the second starts with the short forms in the header's own formatting, repeating;
    same column widths, nothing but an empty paragraph between them; the PDF printed from the split file."""
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Pt

    template = _wide_template(tmp_path / "wide.docx")
    src = from_docx(str(template), str(tmp_path / "figures")).document
    [table] = [b for b in src.blocks if b.kind == "table"]
    table.rows[0][8] = "{{ deadline }}"
    mapping = dict(zip(table.columns, SHORT, strict=True))
    result = fill_docx(str(template), to_markdown(src), {"deadline": "2027-01-31"}, str(tmp_path / "out" / "wide.docx"),
                       pdf=True, abridged_headers=mapping)
    assert result.issues == [] and result.pdf == str(tmp_path / "out" / "wide.pdf") and result.replaced == 1
    assert result.abridged == [table.columns[0]] and result.pages and result.pages >= 3
    assert pathlib.Path(result.pdf).stat().st_size > 1000
    out = docx.Document(result.docx)
    first, rest = out.tables  # two where there was one, the rows going on where page one ended
    assert 2 < len(first.rows) < 81 and len(first.rows) + len(rest.rows) == 82
    assert [c.text for c in first.rows[0].cells] == list(table.columns) and not _is_header(first)
    assert [c.text for c in rest.rows[0].cells] == SHORT and _is_header(rest)
    assert rest.cell(1, 0).text == f"r{len(first.rows)}c0" and first.cell(1, 8).text == "2027-01-31"
    assert _grid(first) == _grid(rest)
    run = next(r for p in rest.rows[0].cells[0].paragraphs for r in p.runs if r.text)
    assert (run.bold, run.font.size, run.font.name) == (True, Pt(9), "Arial")
    body = out.element.body
    between = list(body)[list(body).index(first._tbl) + 1:list(body).index(rest._tbl)]
    assert between and all(el.tag == qn("w:p") and not "".join(el.itertext()).strip() for el in between)


@pytest.mark.skipif(not WINWORD.is_file() or not (QF17 / "qf-17-safety-plan.md").is_file(),
                    reason="Word and asi_standards (private, beside this repo) are both needed")
def test_qf_17_abridges_its_plan_table(tmp_path):
    """The owner's QF-17 (0.8.0): its ten wordy headings verbatim on the plan table's first page, abridged on
    every later one; the split carries the first part's widths, the file keeps its read-only recommendation
    and its final mark, the document is shorter."""
    import docx
    from docx.oxml.ns import qn

    from asi_gutenberg import parameters

    source = (QF17 / "qf-17-safety-plan.md").read_text(encoding="utf-8")
    template = next((QF17 / "source").glob("*.docx"))
    doc = from_markdown(source)
    plan_source = next(b for b in doc.blocks if b.kind == "table" and b.columns[0].startswith("PDP Phases"))
    mapping = dict(zip(plan_source.columns, SHORT, strict=True))  # the headings as the imported source has them
    values = {n: f"<{n}>" for n in parameters(doc)}
    values["vehicle_type"] = "Wheel Loader"
    result = fill_docx(str(template), source, values, str(tmp_path / "qf-17.docx"), pdf=True, abridged_headers=mapping)
    assert result.issues == [] and result.abridged == ["PDP Phases"] and result.pdf == str(tmp_path / "qf-17.pdf")
    assert result.pages and result.pages < 19  # 19 with the verbatim header on every page
    out = docx.Document(result.docx)
    assert len(out.tables) == 4
    first, rest = out.tables[2], out.tables[3]
    assert len(first.rows) + len(rest.rows) == 51 and len(first.rows) >= 2
    assert [c.text.strip("\xa0 ") for c in first.rows[0].cells][:2] == ["PDP Phases", "Safety\xa0Lifecycle Phase"]
    assert not _is_header(first) and [c.text for c in rest.rows[0].cells] == SHORT and _is_header(rest)
    assert _grid(first) == _grid(rest) and first.style.name == rest.style.name == "Grid Table 6 Colorful"
    wp = out.settings.element.find(qn("w:writeProtection"))
    assert wp is not None and wp.get(qn("w:recommended")) == "1"
    import zipfile

    custom = zipfile.ZipFile(result.docx).read("docProps/custom.xml").decode("utf-8")
    assert '<vt:bool>true</vt:bool>' in custom.split('name="_MarkAsFinal"', 1)[1][:80]


# -- links (0.9.0) -------------------------------------------------------------------------------------
BLUE = THEME["link_bg"].lstrip("#")


def _links_of(cell) -> list[tuple[str, str, str | None, str | None, str | None]]:
    """Every hyperlink of a cell: (target, text, character style, colour, shading fill)."""
    from docx.oxml.ns import qn

    def attr(rpr, tag: str, name: str) -> str | None:
        el = rpr.find(qn(tag)) if rpr is not None else None
        return el.get(qn(name)) if el is not None else None

    out = []
    for h in cell._tc.iter(qn("w:hyperlink")):
        target = cell.part.rels[h.get(qn("r:id"))].target_ref
        text = "".join(t.text or "" for t in h.iter(qn("w:t")))
        rpr = h.find(qn("w:r") + "/" + qn("w:rPr"))
        out.append((target, text, attr(rpr, "w:rStyle", "w:val"), attr(rpr, "w:color", "w:val"),
                    attr(rpr, "w:shd", "w:fill")))
    return out


def _linked_source(template: pathlib.Path, tmp_path: pathlib.Path) -> str:
    """The template's source with the plan table's deliverable a link to the row's url, the role cell
    carrying a link after other text, and the second deliverable's url left for the fill to miss."""
    src = from_docx(str(template), str(tmp_path / "figures")).document
    [_control, plan] = [b for b in src.blocks if b.kind == "table"]
    plan.rows[0][1] = "[Safety Plan]({{ row_1_url }})"
    plan.rows[0][2] = "{{ status }}"
    plan.rows[1][1] = "Role A: [{{ a }}]({{ row_2_url }})\nRole B: {{ b }}"
    return to_markdown(src)


LINK_VALUES = {"row_1_url": "https://www.microsoft.com/", "row_2_url": "https://jama.example.com/items/2554",
               "status": "Pending", "a": "Ann", "b": "Ben"}


def test_a_link_in_a_cell_becomes_a_blue_hyperlink(tmp_path):
    """0.9.0, the owner: "turn the name of the safety deliverable into a clickable URL and highlight it blue
    instead of green so it pops". The source writes ``[Safety Plan]({{ row_1_url }})``; filled, the cell is a
    ``w:hyperlink`` to the url, its run in the Hyperlink style, Word's blue, underlined, shaded the theme's
    pale blue and never the mark's green; a link inside a cell with other text links its own span only."""
    import docx
    from docx.oxml.ns import qn

    template = _template(tmp_path / "widget.docx")
    source = _linked_source(template, tmp_path)
    result = fill_docx(str(template), source, LINK_VALUES, str(tmp_path / "linked.docx"), mark=True)
    assert result.unfilled == [] and result.issues == []
    out = docx.Document(result.docx)
    plan = out.tables[1]
    deliverable = plan.cell(1, 1)
    assert deliverable.text == "Safety Plan"
    assert _links_of(deliverable) == [("https://www.microsoft.com/", "Safety Plan", "Hyperlink", "0563C1", BLUE)]
    [h] = list(deliverable._tc.iter(qn("w:hyperlink")))
    assert out.part.rels[h.get(qn("r:id"))].is_external
    u = h.find(qn("w:r") + "/" + qn("w:rPr") + "/" + qn("w:u"))
    assert u is not None and u.get(qn("w:val")) == "single"
    assert not any(s.get(qn("w:fill")) == GREEN for s in deliverable._tc.iter(qn("w:shd")))  # blue, not green
    assert "Hyperlink" in [s.name for s in out.styles]  # the style added: python-docx's default has none
    # the status cell beside it is still a marked value, green
    [status] = [r for p in plan.cell(1, 2).paragraphs for r in p.runs if r.text]
    assert status.text == "Pending" and status._r.find(qn("w:rPr") + "/" + qn("w:shd")).get(qn("w:fill")) == GREEN
    # a link after other text, in a cell written line by line: the prefix stays plain, the link is its span
    roles = plan.cell(2, 1)
    assert _texts(roles) == ["Role A: Ann", "Role B: Ben"]
    assert _links_of(roles) == [("https://jama.example.com/items/2554", "Ann", "Hyperlink", "0563C1", BLUE)]
    assert [r.text for r in roles.paragraphs[0].runs] == ["Role A", ": "]  # the template's run, then the new text
    assert all(r._r.find(qn("w:rPr") + "/" + qn("w:shd")) is None for r in roles.paragraphs[0].runs)
    # read back through the importer, the file says the plain text of the filled source
    again = from_docx(result.docx, str(tmp_path / "figures2")).document
    assert [b for b in again.blocks if b.kind == "table"][1].rows[0][1] == "Safety Plan"


def test_a_link_whose_url_is_unfilled_is_its_plain_text(tmp_path):
    """A link whose url still holds its ``{{ placeholder }}`` (or is empty) collapses to the text: no
    hyperlink, the cell's text as the template had it; the url is listed as unfilled. Since 0.10.0 that
    text is shaded the attention pinkish red with ``mark`` (the owner: a blank link is a gap to see), and
    not at all without it."""
    import docx
    from docx.oxml.ns import qn

    template = _template(tmp_path / "widget.docx")
    source = _linked_source(template, tmp_path)
    values = {"status": "Pending", "a": "Ann", "b": "Ben"}
    result = fill_docx(str(template), source, values, str(tmp_path / "unlinked.docx"), mark=True)
    assert result.unfilled == ["row_1_url", "row_2_url"]
    out = docx.Document(result.docx)
    plan = out.tables[1]
    assert plan.cell(1, 1).text == "Safety Plan" and _links_of(plan.cell(1, 1)) == []
    assert _fills_of(plan.cell(1, 1)) == [("Safety Plan", PINK)]  # a gap: pinkish red, no link
    assert _texts(plan.cell(2, 1)) == ["Role A: Ann", "Role B: Ben"] and _links_of(plan.cell(2, 1)) == []
    [ann] = [r for r in plan.cell(2, 1).paragraphs[0].runs if r.text == "Ann"]
    assert ann._r.find(qn("w:rPr") + "/" + qn("w:shd")).get(qn("w:fill")) == PINK  # the link's gap over the mark
    src = from_markdown(source)
    [_control, plan_source] = [b for b in src.blocks if b.kind == "table"]
    plan_source.rows[0][1] = "[Safety Plan]()"
    empty = fill_docx(str(template), to_markdown(src), {**values, "row_2_url": ""}, str(tmp_path / "empty.docx"),
                      mark=True)
    plan = docx.Document(empty.docx).tables[1]
    assert plan.cell(1, 1).text == "Safety Plan" and _links_of(plan.cell(1, 1)) == []
    assert _fills_of(plan.cell(1, 1)) == [("Safety Plan", PINK)] and _links_of(plan.cell(2, 1)) == []
    exact = fill_docx(str(template), to_markdown(src), {**values, "row_2_url": ""}, str(tmp_path / "exact.docx"))
    plan = docx.Document(exact.docx).tables[1]  # without mark: the template exactly, nothing shaded
    assert plan.cell(1, 1).text == "Safety Plan" and not any(True for _ in plan.cell(1, 1)._tc.iter(qn("w:shd")))


# -- attention (0.10.0) --------------------------------------------------------------------------------
PINK = THEME["warn_bg"].lstrip("#")
DARK_RED = THEME["warn_fg"].lstrip("#")
MS = "https://www.microsoft.com/"


def _fills_of(cell) -> list[tuple[str, str | None]]:
    """Every run of a cell with text, outside and inside its hyperlinks: (text, shading fill)."""
    from docx.oxml.ns import qn

    out = []
    for r in cell._tc.iter(qn("w:r")):
        text = "".join(t.text or "" for t in r.findall(qn("w:t")))
        if not text:
            continue
        shd = r.find(qn("w:rPr") + "/" + qn("w:shd"))
        out.append((text, shd.get(qn("w:fill")) if shd is not None else None))
    return out


def test_attention_values_and_links_are_pinkish_red(tmp_path):
    """0.10.0, the owner: "when the link is microsoft.com because of default, or blank, highlight in pinkish
    red instead of blue; and when no Jira ticket is attached also highlight as pinkish red." ``warn`` names
    the parameters whose values are gaps: shaded ``warn_bg``, never green; ``warn_urls`` the placeholder
    urls: a link to one is still a hyperlink to it, in the Hyperlink style, but dark red over ``warn_bg``
    instead of blue; every other link stays blue and every other value green."""
    import docx
    from docx.oxml.ns import qn

    template = _template(tmp_path / "widget.docx")
    source = _linked_source(template, tmp_path)
    values = {**LINK_VALUES, "status": "no ticket attached"}
    result = fill_docx(str(template), source, values, str(tmp_path / "attention.docx"), mark=True,
                       warn=["status"], warn_urls=[MS])
    assert result.unfilled == [] and result.issues == []
    plan = docx.Document(result.docx).tables[1]
    deliverable = plan.cell(1, 1)
    assert deliverable.text == "Safety Plan"
    assert _links_of(deliverable) == [(MS, "Safety Plan", "Hyperlink", DARK_RED, PINK)]  # still a link, not blue
    [h] = list(deliverable._tc.iter(qn("w:hyperlink")))
    assert h.find(qn("w:r") + "/" + qn("w:rPr") + "/" + qn("w:u")).get(qn("w:val")) == "single"
    assert _fills_of(plan.cell(1, 2)) == [("no ticket attached", PINK)]  # the gap, never green
    assert _links_of(plan.cell(2, 1)) == [("https://jama.example.com/items/2554", "Ann", "Hyperlink", "0563C1", BLUE)]
    [ben] = [r for r in plan.cell(2, 1).paragraphs[1].runs if r.text == "Ben"]
    assert ben._r.find(qn("w:rPr") + "/" + qn("w:shd")).get(qn("w:fill")) == GREEN  # a value: still green
    assert not any(s.get(qn("w:fill")) == GREEN for s in deliverable._tc.iter(qn("w:shd")))
    again = from_docx(result.docx, str(tmp_path / "figures2")).document  # the file reads back plain
    assert [b for b in again.blocks if b.kind == "table"][1].rows[0][1:] == ["Safety Plan", "no ticket attached"]
    # without mark the file is the template exactly: the value plain, the placeholder link blue
    exact = fill_docx(str(template), source, values, str(tmp_path / "exact.docx"), warn=["status"], warn_urls=[MS])
    plan = docx.Document(exact.docx).tables[1]
    assert _links_of(plan.cell(1, 1)) == [(MS, "Safety Plan", "Hyperlink", "0563C1", BLUE)]
    assert _fills_of(plan.cell(1, 2)) == [("no ticket attached", None)]
    # an empty warn value stays unmarked, as any other
    blank = fill_docx(str(template), source, {**values, "status": ""}, str(tmp_path / "blank.docx"), mark=True,
                      warn=["status"])
    assert _fills_of(docx.Document(blank.docx).tables[1].cell(1, 2)) == []


def test_the_command_line_takes_the_attention_flags(tmp_path):
    """``fill --warn NAME --warn-url URL`` (repeatable), with ``--mark``."""
    import docx

    (tmp_path / "source").mkdir()
    template = _template(tmp_path / "source" / "widget.docx")
    md = tmp_path / "widget.md"
    md.write_text(_linked_source(template, tmp_path), encoding="utf-8")
    (tmp_path / "values.json").write_text(json.dumps({**LINK_VALUES, "status": "no ticket attached"}),
                                          encoding="utf-8")
    run = subprocess.run([sys.executable, "-m", "asi_gutenberg", "fill", str(md), str(tmp_path / "out" / "w.docx"),
                          "--params-file", str(tmp_path / "values.json"), "--mark", "--warn", "status",
                          "--warn", "owner", "--warn-url", MS, "--warn-url", "https://nowhere.example/"],
                         capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 0, run.stderr
    payload = json.loads(run.stdout)
    assert payload["ok"] and payload["unfilled"] == []
    plan = docx.Document(payload["docx"]).tables[1]
    assert _links_of(plan.cell(1, 1)) == [(MS, "Safety Plan", "Hyperlink", DARK_RED, PINK)]
    assert _fills_of(plan.cell(1, 2)) == [("no ticket attached", PINK)]
    assert _links_of(plan.cell(2, 1))[0][4] == BLUE


@pytest.mark.skipif(not WINWORD.is_file(), reason="Word is not installed here")
def test_word_prints_a_linked_cell(tmp_path):
    template = _template(tmp_path / "widget.docx")
    result = fill_docx(str(template), _linked_source(template, tmp_path), LINK_VALUES,
                       str(tmp_path / "linked.docx"), mark=True, pdf=True)
    assert result.issues == [] and result.pdf == str(tmp_path / "linked.pdf")
    assert pathlib.Path(result.pdf).stat().st_size > 1000


def _winword_pids() -> set[int]:
    """The WINWORD.EXE processes running now, by id."""
    out = subprocess.run(["tasklist", "/fi", "imagename eq WINWORD.EXE", "/fo", "csv", "/nh"],
                         capture_output=True, text=True).stdout
    return {int(row.split('","')[1]) for row in out.splitlines() if row.startswith('"WINWORD')}


@pytest.mark.skipif(not WINWORD.is_file(), reason="Word is not installed here")
def test_two_consecutive_word_prints_both_succeed(tmp_path):
    """0.9.1: a Word the last session quit could still be shutting down when the next one started, its class
    factory still registered, so the next activation landed on it and died with it (one Word test in three
    failed with an RPC error). A session now waits for the Word it started to be gone before it ends, so two
    prints in a row both succeed and no Word of ours is running between them."""
    template = _template(tmp_path / "widget.docx")
    source = _source(template, tmp_path)
    before = _winword_pids()
    for name in ("first", "second"):
        result = fill_docx(str(template), source, VALUES, str(tmp_path / f"{name}.docx"), pdf=True)
        assert result.issues == [] and result.pdf == str(tmp_path / f"{name}.pdf"), name
        assert pathlib.Path(result.pdf).stat().st_size > 1000
        assert _winword_pids() <= before, f"{name}: the Word the session started is still running"
